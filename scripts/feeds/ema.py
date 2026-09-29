#!/usr/bin/env python3
"""EU catalogue feed: EMA EPAR "Medicines" report.

Downloads the Medicines report (XLSX), converts it to the `;` CSV that
EmaEparParser reads and publishes data/eu/ema-epar-<yyyymm>.zip with it
as `ema-epar.csv` (docs/analysis/ANALYSIS-CATALOGUE-REMOTE-FEEDS-EU-ES-FR.md
§3.1).

Conversion rules:
  * sheet `Medicine`, or the first one;
  * the header is the first row whose first cell is `Category` (the
    report puts 8 metadata rows above it today); trailing empty header
    cells are dropped (the sheet declares 1,024 columns for 39 real
    ones) and every row is cut to the header width;
  * fully empty rows are skipped;
  * cells: empty for None; CR, LF and TAB become spaces and runs of
    spaces collapse (the parser reads one record per physical line);
    integral numbers lose the `.0`; dates are written as yyyy-mm-dd;
  * `;` delimiter, UTF-8 without BOM, `\\n` line ends, fields quoted
    only when needed (many cells hold `;`-separated lists).
The CSV is then re-read and checked before anything is published.
"""

import csv
import datetime
import re
import sys
import tempfile
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
from bs4 import BeautifulSoup

import common

SOURCE_URL = "https://www.ema.europa.eu/en/documents/report/medicines-output-medicines-report_en.xlsx"
LANDING_URL = "https://www.ema.europa.eu/en/medicines/download-medicine-data"
REPORT_NAME = "medicines-output-medicines-report_en.xlsx"
COUNTRY = "EU"
PREFIX = "ema-epar"
DATA_DIR = Path("data/eu")
ENTRY_NAME = "ema-epar.csv"
HUMAN_ROWS_KEY = "ema-epar.csv (Human)"

SHEET_NAME = "Medicine"
HEADER_FIRST_CELL = "Category"

# The columns EmaEparParser requires (matched case-insensitively).
REQUIRED_COLUMNS = [
    "Category",
    "Name of medicine",
    "EMA product number",
    "Medicine status",
    "Active substance",
    "ATC code (human)",
    "Marketing authorisation developer / applicant / holder",
    "Medicine URL",
]

# Absolute floors, below the real sizes (2,746 rows, 2,351 of them
# Human, on 2026-09-29) but far above a truncated download.
MIN_ROWS = 2_000
MIN_HUMAN_ROWS = 1_800

ZIP_MAGIC = b"PK\x03\x04"
_WHITESPACE_RUN = re.compile(r"[\r\n\t ]+")


def clean_cell(value):
    """One spreadsheet value as CSV text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _WHITESPACE_RUN.sub(" ", str(value)).strip()


def convert_rows(rows):
    """Header and data rows (lists of strings) from raw sheet rows."""
    rows = iter(rows)
    header = None
    for row in rows:
        cells = [clean_cell(value) for value in row]
        if cells and cells[0] == HEADER_FIRST_CELL:
            header = cells
            break
    if header is None:
        raise common.FeedError(f"no row starts with '{HEADER_FIRST_CELL}'")
    while header and not header[-1]:
        header.pop()

    width = len(header)
    data = []
    for row in rows:
        cells = [clean_cell(value) for value in row][:width]
        if not any(cells):
            continue
        cells.extend([""] * (width - len(cells)))
        data.append(cells)
    return header, data


def read_workbook(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.worksheets[0]
        return convert_rows(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def write_csv(path, header, data):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(data)


def convert(xlsx_path, csv_path):
    header, data = read_workbook(xlsx_path)
    write_csv(csv_path, header, data)


def validate(csv_path, previous):
    """Re-read the CSV as the parser will; return the manifest row counts."""
    with open(csv_path, encoding="utf-8", newline="") as fh:
        text = fh.read()
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()

    # One record per physical line, as the parser reads it: each line is
    # parsed on its own (a reader over all lines would silently join a
    # quoted field that spans two) and must give the header's width.
    if not lines:
        raise common.FeedError(f"{ENTRY_NAME}: empty")
    parsed = []
    for number, line in enumerate(lines, start=1):
        try:
            row = next(csv.reader([line], delimiter=";", strict=True), [])
        except csv.Error as error:
            raise common.FeedError(f"{ENTRY_NAME}: line {number} is not a complete record ({error})") from error
        if parsed and len(row) != len(parsed[0]):
            raise common.FeedError(
                f"{ENTRY_NAME}: line {number} has {len(row)} columns, {len(parsed[0])} expected"
            )
        parsed.append(row)
    header = parsed[0]

    folded = [name.casefold() for name in header]
    missing = [name for name in REQUIRED_COLUMNS if name.casefold() not in folded]
    if missing:
        raise common.FeedError(f"{ENTRY_NAME}: missing columns {missing}")

    category = folded.index(HEADER_FIRST_CELL.casefold())
    data = parsed[1:]
    human = sum(1 for row in data if row[category].casefold() == "human")

    return {
        ENTRY_NAME: common.check_rows(ENTRY_NAME, len(data), MIN_ROWS, previous.get(ENTRY_NAME)),
        HUMAN_ROWS_KEY: common.check_rows(HUMAN_ROWS_KEY, human, MIN_HUMAN_ROWS, previous.get(HUMAN_ROWS_KEY)),
    }


def find_report_link(html, page_url=LANDING_URL):
    soup = BeautifulSoup(html, "html.parser")
    for anchor in soup.find_all("a", href=True):
        if anchor["href"].split("?", 1)[0].endswith(REPORT_NAME):
            return urljoin(page_url, anchor["href"])
    return None


def check_payload(data):
    if common.looks_like_html(data):
        raise common.FeedError(f"{REPORT_NAME}: the server returned an HTML page")
    if not data.startswith(ZIP_MAGIC):
        raise common.FeedError(f"{REPORT_NAME}: not an XLSX (ZIP) container")


def download(session):
    # HEAD reports a length of 0 for this file (a cached response), so
    # the size is only known from the body.
    headers = common.browser_headers(LANDING_URL)
    try:
        print(f"Download {SOURCE_URL}")
        response = session.get(SOURCE_URL, headers=headers, timeout=120)
        response.raise_for_status()
        check_payload(response.content)
        return response.content
    except (common.FeedError, OSError) as error:
        # requests' exceptions derive from OSError.
        print(f"Direct download failed ({error}); looking for the link on {LANDING_URL}")

    page = session.get(LANDING_URL, headers=headers, timeout=60)
    page.raise_for_status()
    url = find_report_link(page.text, page.url or LANDING_URL)
    if url is None:
        raise common.FeedError(f"no link to {REPORT_NAME} on {LANDING_URL}")
    print(f"Download {url}")
    response = session.get(url, headers=headers, timeout=120)
    response.raise_for_status()
    check_payload(response.content)
    return response.content


def main():
    run = common.run_time()
    version = common.version_for(run)
    if common.already_published(DATA_DIR, version, common.force_refresh()):
        print(f"Version {version} is already published; nothing to do.")
        return 0

    payload = download(common.session())
    print(f"Downloaded {len(payload):,} bytes")

    with tempfile.TemporaryDirectory() as work:
        workbook = Path(work) / REPORT_NAME
        workbook.write_bytes(payload)
        csv_path = Path(work) / ENTRY_NAME
        convert(workbook, csv_path)
        rows = validate(csv_path, common.previous_rows(DATA_DIR))

        archive = common.write_zip(Path(work) / f"{PREFIX}-{version}.zip", {ENTRY_NAME: csv_path})
        common.publish(DATA_DIR, COUNTRY, version, archive, 1, rows, run)
    return 0


if __name__ == "__main__":
    sys.exit(main())
